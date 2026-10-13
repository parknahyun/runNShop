"""
ADM-ORDER-01 (어드민 주문 현황) 와이어프레임 초안 생성 스크립트
화면정의서_작성_표준안.md §5 규칙에 따라 PPT로 작성 → PDF로 내보내는 흐름의 PPT 산출.
소스: ADM-ORDER-01_어드민-주문목록.md 의 구성요소 인덱스(R1~R6, R0-S1~S3)를 그대로 시각화한다.

영역/속성/기능 박스 자체는 초안(첫 버전) 그대로 중립 스타일(옅은 배경 또는 테두리만)로 그린다.
코드 라벨 텍스트박스([R1], [P1], [F1] 등)만 종류별로 배경색을 구분한다:
- 영역(Region, R) 라벨: 연파란색
- 속성(Property, P) 라벨: 연녹색
- 기능(Function, F) 라벨: 연노란색

실행: python3 make_adm_order_01_wireframe.py
출력: ../ADM-ORDER-01.pptx (같은 폴더의 상위, 화면정의서_샘플/)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "ADM-ORDER-01.pptx")

# 색상 (박스 자체는 초안 스타일 — 중립)
C_BORDER = RGBColor(0x33, 0x41, 0x55)
C_HEADER_BG = RGBColor(0x1E, 0x29, 0x3B)
C_HEADER_TXT = RGBColor(0xF1, 0xF5, 0xF9)
C_LABEL = RGBColor(0x94, 0xA3, 0xB8)
C_BODY_TXT = RGBColor(0x33, 0x41, 0x55)
C_LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)

# 코드 라벨([R1]/[P1]/[F1]) 전용 색상 — 표준안 4장 코드 체계와 매핑
C_REGION_FILL = RGBColor(0xDB, 0xEA, 0xFE)    # 영역(R) 라벨 — 연파란색
C_PROPERTY_FILL = RGBColor(0xDC, 0xFC, 0xE7)  # 속성(P) 라벨 — 연녹색
C_FUNCTION_FILL = RGBColor(0xFE, 0xF9, 0xC3)  # 기능(F) 라벨 — 연노란색
C_REGION_BORDER = RGBColor(0x93, 0xC5, 0xFD)
C_PROPERTY_BORDER = RGBColor(0x86, 0xEF, 0xAC)
C_FUNCTION_BORDER = RGBColor(0xFD, 0xE0, 0x47)

C_BADGE = {
    "접수": RGBColor(0x94, 0xA3, 0xB8),
    "결제완료": RGBColor(0x0E, 0xA5, 0xE9),
    "배송중": RGBColor(0xF5, 0x9E, 0x0B),
    "구매확정": RGBColor(0x22, 0xC5, 0x5E),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]


def add_textbox(slide, left, top, width, height, text, size=11, bold=False,
                 color=C_BODY_TXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
                 font="Malgun Gothic"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def add_rect(slide, left, top, width, height, fill=None, line_color=C_BORDER,
             line_width=Pt(1), dashed=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line_color
    shp.line.width = line_width
    if dashed:
        shp.line.dash_style = 2  # MSO_LINE_DASH_STYLE.DASH approx
    shp.shadow.inherit = False
    return shp


def _code_kind(code_text):
    """코드 텍스트(예: "[R1]", "[P1 x8]", "[F1] 행 클릭→상세")의 접두 문자(R/P/F)로 종류 판별"""
    inner = code_text.lstrip("[")
    if inner.startswith("P"):
        return "P"
    if inner.startswith("F"):
        return "F"
    if inner.startswith("R"):
        return "R"
    return None


def add_code_label(slide, left, top, width, height, code_text, kind=None, size=9, align=PP_ALIGN.LEFT):
    """영역/속성/기능 코드 라벨 텍스트박스 — 종류별 배경색 적용
    R(영역)=연파란색, P(속성)=연녹색, F(기능)=연노란색. kind 미지정 시 code_text로 자동 판별."""
    kind = kind or _code_kind(code_text) or "R"
    fill, border = {
        "R": (C_REGION_FILL, C_REGION_BORDER),
        "P": (C_PROPERTY_FILL, C_PROPERTY_BORDER),
        "F": (C_FUNCTION_FILL, C_FUNCTION_BORDER),
    }[kind]
    add_rect(slide, left, top, width, height, fill=fill, line_color=border, line_width=Pt(0.75))
    add_textbox(slide, left, top, width, height, code_text,
                size=size, bold=True, color=C_BODY_TXT, align=align, anchor=MSO_ANCHOR.MIDDLE)


def add_region_label(slide, left, top, code_text):
    """영역 코드 라벨(예: [R1])을 우측 상단 바깥쪽에 작게 표기 — 연파란색 배경"""
    add_code_label(slide, left, top, Inches(0.9), Inches(0.25), code_text, kind="R")


def add_marker(slide, left, top, code_text, w=Inches(0.55)):
    """개별 속성/기능 마커(예: [P1], [F1]) — 코드 접두에 따라 연녹색/연노란색 배경"""
    add_code_label(slide, left, top, w, Inches(0.2), code_text, size=8)


# ---------------------------------------------------------------
# Slide 1: 정상 조회 상태 (R1~R6 전부 노출)
# ---------------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)

add_textbox(slide, Inches(0.3), Inches(0.05), Inches(9), Inches(0.3),
            "ADM-ORDER-01 — 어드민 주문 현황 (OrderHistory) · 목록형 · 상태: 조회 후(R0-S1~S3 해당 없음)",
            size=11, bold=True, color=C_HEADER_BG)

top = Inches(0.45)
left = Inches(0.4)
content_w = Inches(12.5)

# R1 PageTitle
r1_h = Inches(0.5)
add_rect(slide, left, top, content_w, r1_h, fill=C_LIGHT_BG)
add_textbox(slide, left + Inches(0.15), top, content_w - Inches(0.3), r1_h,
            "주문 현황", size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
add_region_label(slide, left + content_w - Inches(0.6), top, "[R1]")
top += r1_h + Inches(0.08)

# R2 FilterBar
r2_h = Inches(0.55)
add_rect(slide, left, top, content_w, r2_h, fill=None)
fx = left + Inches(0.15)
fw = Inches(2.6)
add_textbox(slide, fx, top, fw, r2_h, "[검색] 주문번호·상품명·고객명", size=10, anchor=MSO_ANCHOR.MIDDLE)
add_marker(slide, fx + fw - Inches(0.1), top, "[P1]")
fx += fw + Inches(0.2)
add_textbox(slide, fx, top, Inches(1.8), r2_h, "[상태 ▾] 전체", size=10, anchor=MSO_ANCHOR.MIDDLE)
add_marker(slide, fx + Inches(1.7), top, "[P2]")
fx += Inches(2.0)
add_textbox(slide, fx, top, Inches(2.6), r2_h, "[기간] ____ ~ ____", size=10, anchor=MSO_ANCHOR.MIDDLE)
add_marker(slide, fx + Inches(2.5), top, "[P3]")
fx += Inches(2.8)
add_rect(slide, fx, top + Inches(0.08), Inches(1.0), Inches(0.38), fill=C_HEADER_BG)
add_textbox(slide, fx, top + Inches(0.08), Inches(1.0), Inches(0.38), "조회",
            size=10, bold=True, color=C_HEADER_TXT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_marker(slide, fx + Inches(1.0), top, "[F1]")
add_region_label(slide, left + content_w - Inches(0.6), top, "[R2]")
top += r2_h + Inches(0.08)

# R3 SummaryBar (8 cards)
r3_h = Inches(0.75)
add_rect(slide, left, top, content_w, r3_h, fill=C_LIGHT_BG)
labels = ["전체", "접수", "결제완료", "준비중", "배송중", "배송완료", "구매확정", "취소"]
counts = ["128", "12", "20", "18", "30", "22", "20", "6"]
card_w = (content_w - Inches(0.3)) / 8
cx = left + Inches(0.15)
for i, (lab, cnt) in enumerate(zip(labels, counts)):
    add_rect(slide, cx, top + Inches(0.08), card_w - Inches(0.06), r3_h - Inches(0.16), fill=None)
    add_textbox(slide, cx, top + Inches(0.1), card_w - Inches(0.06), Inches(0.25), lab,
                size=9, color=C_LABEL, align=PP_ALIGN.CENTER)
    add_textbox(slide, cx, top + Inches(0.35), card_w - Inches(0.06), Inches(0.3), cnt,
                size=13, bold=True, align=PP_ALIGN.CENTER)
    cx += card_w
add_marker(slide, left + Inches(0.15), top, "[P1 x8]")
add_region_label(slide, left + content_w - Inches(0.6), top, "[R3]")
top += r3_h + Inches(0.08)

# R4 Toolbar
r4_h = Inches(0.42)
add_rect(slide, left, top, content_w, r4_h, fill=None)
add_rect(slide, left + content_w - Inches(1.6), top + Inches(0.05), Inches(1.45), Inches(0.32), fill=None)
add_textbox(slide, left + content_w - Inches(1.6), top + Inches(0.05), Inches(1.45), Inches(0.32),
            "엑셀 다운로드", size=9, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_marker(slide, left + content_w - Inches(0.2), top, "[F1]")
add_region_label(slide, left + content_w - Inches(0.6) - Inches(1.7), top, "[R4]")
top += r4_h + Inches(0.05)

# R5 OrderTable
r5_h = Inches(2.6)
add_rect(slide, left, top, content_w, r5_h, fill=None)
cols = ["주문번호", "상품명", "주문자", "입점사", "상태", "결제금액", "주문일시"]
col_w = (content_w) / len(cols)
hx = left
header_h = Inches(0.35)
add_rect(slide, left, top, content_w, header_h, fill=RGBColor(0xE2, 0xE8, 0xF0))
for i, c in enumerate(cols):
    add_textbox(slide, hx, top, col_w, header_h, c, size=9, bold=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    hx += col_w
p_codes = ["P1", "P2", "P3", "P4", "P5", "P6", "P7"]
hx = left
for i, pc in enumerate(p_codes):
    add_marker(slide, hx + col_w - Inches(0.45), top + Inches(0.35), f"[{pc}]", w=Inches(0.45))
    hx += col_w

row_h = Inches(0.35)
sample_rows = [
    ("ORD-20260812-00931", "나이키 페가수스 41", "홍*동", "나이키코리아", "배송중", "152,000원", "2026-08-12 14:02"),
    ("ORD-20260812-00930", "삼다수 2L x 12", "김*희", "제주삼다수", "결제완료", "18,900원", "2026-08-12 13:47"),
    ("ORD-20260812-00929", "무선 이어폰 Pro", "이*수", "사운드랩", "구매확정", "129,000원", "2026-08-12 11:20"),
    ("ORD-20260811-00928", "런닝 크루넥 티셔츠", "박*영", "런앤샵", "접수", "39,000원", "2026-08-11 22:05"),
]
ry = top + header_h
for row in sample_rows:
    for i, val in enumerate(row):
        if i == 4 and val in C_BADGE:
            add_rect(slide, left + col_w * i + Inches(0.15), ry + Inches(0.05),
                      Inches(0.9), Inches(0.25), fill=C_BADGE[val])
            add_textbox(slide, left + col_w * i + Inches(0.15), ry + Inches(0.05),
                        Inches(0.9), Inches(0.25), val, size=8, bold=True,
                        color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            add_textbox(slide, left + col_w * i, ry, col_w, row_h, val, size=9,
                        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ry += row_h
add_marker(slide, left + Inches(0.05), ry + Inches(0.02), "[F1] 행 클릭→상세", w=Inches(1.8))
add_region_label(slide, left + content_w - Inches(0.6), top, "[R5]")
top += r5_h + Inches(0.08)

# R6 Pagination
r6_h = Inches(0.4)
add_rect(slide, left, top, content_w, r6_h, fill=None)
add_textbox(slide, left + Inches(0.15), top, Inches(4), r6_h,
            "‹ 1 2 3 4 5 ›     페이지당 [20 ▾]건", size=10, anchor=MSO_ANCHOR.MIDDLE)
add_marker(slide, left + Inches(2.3), top, "[F1]")
add_marker(slide, left + Inches(4.2), top, "[F2]")
add_region_label(slide, left + content_w - Inches(0.6), top, "[R6]")


def draw_frame_skeleton(slide, top, left, content_w, show_r3_r6, r5_message=None):
    """R1(PageTitle)~R2(FilterBar)는 항상 동일, R3~R6은 상태에 따라 렌더링/대체.
    show_r3_r6=True: 정상(R3~R6 전부) / False: R3·R6 미표시, R5 자리에 r5_message 표시.
    R4(Toolbar)는 두 경우 모두 노출(실제 구현과 동일)."""
    y = top

    # R1
    r1_h = Inches(0.45)
    add_rect(slide, left, y, content_w, r1_h, fill=C_LIGHT_BG)
    add_textbox(slide, left + Inches(0.15), y, content_w - Inches(0.3), r1_h,
                "주문 목록", size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_region_label(slide, left + content_w - Inches(0.6), y, "[R1]")
    y += r1_h + Inches(0.06)

    # R2 FilterBar (항상 노출)
    r2_h = Inches(0.45)
    add_rect(slide, left, y, content_w, r2_h, fill=None)
    add_textbox(slide, left + Inches(0.15), y, Inches(3.0), r2_h,
                "[검색] · [상태▾] · [기간] · [조회]", size=9, anchor=MSO_ANCHOR.MIDDLE)
    add_region_label(slide, left + content_w - Inches(0.6), y, "[R2]")
    y += r2_h + Inches(0.06)

    if show_r3_r6:
        # R3 SummaryBar
        r3_h = Inches(0.55)
        add_rect(slide, left, y, content_w, r3_h, fill=C_LIGHT_BG)
        add_textbox(slide, left + Inches(0.15), y, content_w - Inches(0.3), r3_h,
                    "전체 128 · 접수 12 · 결제완료 20 · 준비중 18 · 배송중 30 · 배송완료 22 · 구매확정 20 · 취소 6",
                    size=9, anchor=MSO_ANCHOR.MIDDLE)
        add_region_label(slide, left + content_w - Inches(0.6), y, "[R3]")
        y += r3_h + Inches(0.06)
    else:
        add_textbox(slide, left + Inches(0.15), y, content_w - Inches(0.3), Inches(0.3),
                    "( R3 SummaryBar — displayed.length === 0 이라 렌더링되지 않음 )",
                    size=8, color=C_LABEL)
        y += Inches(0.06)

    # R4 Toolbar (항상 노출)
    r4_h = Inches(0.35)
    add_rect(slide, left, y, content_w, r4_h, fill=None)
    add_textbox(slide, left + content_w - Inches(1.5), y + Inches(0.03), Inches(1.35), r4_h - Inches(0.06),
                "엑셀 다운로드", size=9, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_region_label(slide, left + content_w - Inches(0.6) - Inches(1.5), y, "[R4]")
    y += r4_h + Inches(0.06)

    if show_r3_r6:
        # R5 OrderTable (정상)
        r5_h = Inches(2.0)
        add_rect(slide, left, y, content_w, r5_h, fill=None)
        add_textbox(slide, left + Inches(0.15), y + Inches(0.1), content_w - Inches(0.3), Inches(0.3),
                    "주문번호 | 상품명 | 주문자 | 입점사 | 상태 | 결제금액 | 주문일시", size=9, bold=True)
        add_textbox(slide, left + Inches(0.15), y + Inches(0.5), content_w - Inches(0.3), Inches(1.3),
                    "ORD-20260812-00931  나이키 페가수스 41  홍*동  나이키코리아  배송중  152,000원  2026-08-12 14:02\n"
                    "ORD-20260812-00930  삼다수 2L x 12         김*희  제주삼다수     결제완료  18,900원  2026-08-12 13:47",
                    size=8)
        add_region_label(slide, left + content_w - Inches(0.6), y, "[R5]")
        y += r5_h + Inches(0.06)

        # R6 Pagination
        r6_h = Inches(0.35)
        add_rect(slide, left, y, content_w, r6_h, fill=None)
        add_textbox(slide, left + Inches(0.15), y, Inches(3), r6_h,
                    "‹ 1 2 3 4 5 ›   페이지당 [20▾]건", size=9, anchor=MSO_ANCHOR.MIDDLE)
        add_region_label(slide, left + content_w - Inches(0.6), y, "[R6]")
        y += r6_h
    else:
        # R5 자리에 상태 메시지가 대신 표시됨 (.table-wrap 컨테이너는 R5와 동일)
        r5_h = Inches(1.2)
        add_rect(slide, left, y, content_w, r5_h, fill=C_LIGHT_BG)
        add_textbox(slide, left, y, content_w, r5_h, r5_message,
                    size=12, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, color=C_LABEL)
        add_region_label(slide, left + content_w - Inches(0.6), y, "[R5]")
        add_textbox(slide, left + Inches(0.15), y + r5_h + Inches(0.06), content_w - Inches(0.3), Inches(0.25),
                    "( R6 Pagination — 위와 동일한 이유로 렌더링되지 않음 )", size=8, color=C_LABEL)
        y += r5_h + Inches(0.35)

    return y


# ---------------------------------------------------------------
# Slide 2: R0-S1 (조회 전) — R5 자리에 안내 문구, R3·R6 미표시, R4는 노출
# ---------------------------------------------------------------
slide2 = prs.slides.add_slide(blank_layout)
add_textbox(slide2, Inches(0.3), Inches(0.05), Inches(11), Inches(0.3),
            "ADM-ORDER-01 — R0-S1 조회 전 (searched=false, R3·R6 미표시 / R4 노출 / R5 자리에 안내문구)",
            size=11, bold=True, color=C_HEADER_BG)
draw_frame_skeleton(slide2, Inches(0.5), Inches(0.4), Inches(12.5),
                     show_r3_r6=False, r5_message="조회 버튼을 눌러 주문 목록을 불러오세요.")

# ---------------------------------------------------------------
# Slide 3: R0-S2 (검색 결과 없음) — 동일 구조, 문구만 다름
# ---------------------------------------------------------------
slide3 = prs.slides.add_slide(blank_layout)
add_textbox(slide3, Inches(0.3), Inches(0.05), Inches(11), Inches(0.3),
            "ADM-ORDER-01 — R0-S2 검색 결과 없음 (searched=true, displayed.length=0)",
            size=11, bold=True, color=C_HEADER_BG)
draw_frame_skeleton(slide3, Inches(0.5), Inches(0.4), Inches(12.5),
                     show_r3_r6=False, r5_message="조건에 맞는 주문이 없습니다.")

# ---------------------------------------------------------------
# Slide 4: R0-S3 (로딩 중) — R2 아래 전체가 로딩 문구로 대체 (R3~R6 전부 미표시)
# ---------------------------------------------------------------
slide4 = prs.slides.add_slide(blank_layout)
add_textbox(slide4, Inches(0.3), Inches(0.05), Inches(11), Inches(0.3),
            "ADM-ORDER-01 — R0-S3 로딩 중 (loading=true, R2 아래 R3~R6 전체가 로딩 문구로 대체)",
            size=11, bold=True, color=C_HEADER_BG)
top4 = Inches(0.5)
left4 = Inches(0.4)
content_w4 = Inches(12.5)
r1_h = Inches(0.45)
add_rect(slide4, left4, top4, content_w4, r1_h, fill=C_LIGHT_BG)
add_textbox(slide4, left4 + Inches(0.15), top4, content_w4 - Inches(0.3), r1_h,
            "주문 목록", size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
add_region_label(slide4, left4 + content_w4 - Inches(0.6), top4, "[R1]")
top4 += r1_h + Inches(0.06)
r2_h = Inches(0.45)
add_rect(slide4, left4, top4, content_w4, r2_h, fill=None)
add_textbox(slide4, left4 + Inches(0.15), top4, Inches(3.0), r2_h,
            "[검색] · [상태▾] · [기간] · [조회]", size=9, anchor=MSO_ANCHOR.MIDDLE)
add_region_label(slide4, left4 + content_w4 - Inches(0.6), top4, "[R2]")
top4 += r2_h + Inches(0.15)
add_rect(slide4, left4, top4, content_w4, Inches(1.5), fill=None)
add_textbox(slide4, left4, top4, content_w4, Inches(1.5), "불러오는 중...",
            size=13, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, color=C_LABEL)
add_textbox(slide4, left4 + Inches(0.15), top4 + Inches(1.55), content_w4 - Inches(0.3), Inches(0.3),
            "( R3~R6 전체가 이 로딩 문구로 대체됨 — .table-wrap 밖의 별도 분기 )", size=8, color=C_LABEL)

os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
prs.save(OUT_PATH)
print(f"Saved: {os.path.abspath(OUT_PATH)}")
